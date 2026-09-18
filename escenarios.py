"""
Sendas de proyección de IPC e IBR por escenario.

El archivo `Escenarios_IPC-IBR_VC.xlsx` trae una hoja por índice con una fila por
fecha y tres columnas de escenario: Alcista, Base y Bajista. Las tres coinciden en
el pasado y divergen hacia adelante, que es justo lo que se espera de una senda.

La lectura de un índice en una fecha es **el último registro publicado con fecha
menor o igual**, sin interpolar. Si la fecha pedida va más allá del último registro
se arrastra ese último valor y se marca como extrapolada, para que el reporte pueda
advertirlo.

La hoja `IB1` del mismo archivo no se usa aquí: la senda diaria real de IBR solo
interviene en el margen de valoración por el atajo de la bvc (ver `ibr.py`).
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

ESCENARIOS = ("Alcista", "Base", "Bajista")
HOJAS = ("IPC", "IBR")
FILA_INICIAL = 2                 # la 1 es el encabezado


class EscenariosFormatError(Exception):
    """El archivo de escenarios no tiene la forma esperada."""


@dataclass
class Senda:
    """Serie de un índice: fechas ordenadas y un valor por escenario."""
    nombre: str
    fechas: list[dt.date] = field(default_factory=list)
    valores: dict[str, list[float]] = field(default_factory=dict)

    @property
    def ultima(self) -> dt.date:
        return self.fechas[-1]

    def vigente(self, fecha: dt.date, escenario: str) -> tuple[float, bool]:
        """Último valor publicado con fecha ≤ la pedida, y si hubo que extrapolar."""
        serie = self.valores[escenario]
        i = -1
        for j, f in enumerate(self.fechas):
            if f <= fecha:
                i = j
            else:
                break
        if i < 0:                                   # antes del primer registro
            return serie[0], True
        return serie[i], fecha > self.fechas[-1]

    def para_json(self) -> dict:
        return {"fechas": [f.isoformat() for f in self.fechas],
                "valores": {e: list(v) for e, v in self.valores.items()}}


@dataclass
class Escenarios:
    sendas: dict[str, Senda] = field(default_factory=dict)
    origen: str = ""

    @property
    def activo(self) -> bool:
        return bool(self.sendas)

    def senda(self, indice: str) -> Senda:
        return self.sendas[indice]

    def para_json(self) -> dict:
        return {"origen": self.origen, "escenarios": list(ESCENARIOS),
                "sendas": {k: s.para_json() for k, s in self.sendas.items()}}

    @classmethod
    def cargar(cls, path: str | Path | None) -> "Escenarios":
        if not path:
            return cls()
        path = Path(path)
        if not path.exists():
            raise EscenariosFormatError(f"No existe el archivo de escenarios: {path}")
        try:
            import openpyxl
        except ImportError:                                   # pragma: no cover
            raise EscenariosFormatError(
                "Para leer el archivo de escenarios hace falta openpyxl: "
                "pip install openpyxl")

        libro = openpyxl.load_workbook(path, data_only=True, read_only=True)
        faltan = [h for h in HOJAS if h not in libro.sheetnames]
        if faltan:
            raise EscenariosFormatError(
                f"{path.name}: faltan las hojas {faltan}. Se esperan {list(HOJAS)}.")

        sendas: dict[str, Senda] = {}
        for hoja in HOJAS:
            s = Senda(hoja, [], {e: [] for e in ESCENARIOS})
            for fila in libro[hoja].iter_rows(min_row=FILA_INICIAL, max_col=4,
                                              values_only=True):
                fecha = fila[0]
                if not isinstance(fecha, dt.datetime):
                    continue
                if any(not isinstance(v, (int, float)) for v in fila[1:4]):
                    continue
                s.fechas.append(fecha.date())
                for e, v in zip(ESCENARIOS, fila[1:4]):
                    s.valores[e].append(float(v) / 100)
            if not s.fechas:
                raise EscenariosFormatError(
                    f"{path.name}, hoja {hoja}: no se leyó ninguna fila. Se espera la "
                    f"fecha en la columna A y {list(ESCENARIOS)} en B, C y D.")
            orden = sorted(range(len(s.fechas)), key=lambda i: s.fechas[i])
            s.fechas = [s.fechas[i] for i in orden]
            s.valores = {e: [v[i] for i in orden] for e, v in s.valores.items()}
            sendas[hoja] = s

        return cls(sendas=sendas, origen=path.name)
